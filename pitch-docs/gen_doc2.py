"""Generate Document 2: Huong dan su dung phan mem Canteen as .docx with screenshots."""
import sys
import io
import os
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from docx import Document
from docx.shared import Pt, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from lxml import etree

BORDER_NONE = (
    '<w:tcBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
    '<w:top w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
    '<w:left w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
    '<w:bottom w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
    '<w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
    '</w:tcBorders>'
)

SCREENSHOT_DIR = os.path.join(os.path.dirname(__file__), '..', 'tools', 'screenshots')


def remove_table_borders(table):
    """Remove all borders from a table."""
    for row in table.rows:
        for cell in row.cells:
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            tcPr.append(etree.fromstring(BORDER_NONE))


def add_run_tnr(paragraph, text, bold=False, italic=False, size=13):
    """Add a run with Times New Roman font."""
    run = paragraph.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    return run


def add_centered(doc, text, bold=False, size=13, space_after=0, space_before=0):
    """Add centered paragraph."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    add_run_tnr(p, text, bold=bold, size=size)
    return p


def add_para(doc, text, bold=False, size=13, indent=0, space_after=6, space_before=0, italic=False):
    """Add a paragraph."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    if indent > 0:
        p.paragraph_format.first_line_indent = Cm(indent)
    add_run_tnr(p, text, bold=bold, italic=italic, size=size)
    return p


def add_h1(doc, text):
    """Add heading level 1."""
    add_para(doc, text, bold=True, size=14, space_before=14, space_after=6)


def add_h2(doc, text):
    """Add heading level 2."""
    add_para(doc, text, bold=True, size=13, space_before=10, space_after=4)


def add_h3(doc, text):
    """Add heading level 3."""
    add_para(doc, text, bold=True, size=13, space_before=8, space_after=4, indent=0.5)


def add_step(doc, text):
    """Add a step paragraph."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Cm(0.5)
    add_run_tnr(p, text)
    return p


def add_bullet(doc, text, indent_cm=1):
    """Add a bullet point."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Cm(indent_cm)
    add_run_tnr(p, '- ' + text)
    return p


def add_note(doc, text):
    """Add a note paragraph."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Cm(0.5)
    add_run_tnr(p, 'L\u01b0u \u00fd: ', bold=True, italic=True)
    add_run_tnr(p, text, italic=True)


def add_image(doc, filename, caption=None, width_inches=5.5):
    """Add a centered image with optional caption below."""
    img_path = os.path.join(SCREENSHOT_DIR, filename)
    if not os.path.exists(img_path):
        add_para(doc, f'[H\u00ecnh: {filename} - kh\u00f4ng t\u00ecm th\u1ea5y]', italic=True, size=11)
        return

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(img_path, width=Inches(width_inches))

    if caption:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(8)
        add_run_tnr(cap, caption, italic=True, size=10)


def add_table_row(table, cells_data, bold=False, size=11):
    """Add a row to a table."""
    row = table.add_row().cells
    for j, val in enumerate(cells_data):
        row[j].paragraphs[0].text = ''
        add_run_tnr(row[j].paragraphs[0], val, bold=bold, size=size)
    return row


def main():
    """Generate the HDSD document with embedded screenshots."""
    doc = Document()

    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(3)
        section.right_margin = Cm(2)

    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(13)

    # ===== COVER PAGE =====
    for _ in range(4):
        doc.add_paragraph()

    add_centered(doc, 'C\u00d4NG TY \u2026\u2026\u2026\u2026\u2026\u2026\u2026\u2026', bold=True, size=14, space_after=6)
    add_centered(doc, '\u2500\u2500\u2500\u2500\u2500 \u2736 \u2500\u2500\u2500\u2500\u2500', size=12, space_after=24)

    for _ in range(2):
        doc.add_paragraph()

    add_centered(doc, 'H\u01af\u1edaNG D\u1eaaN S\u1eda D\u1ee4NG', bold=True, size=18, space_after=6)
    add_centered(doc, 'PH\u1ea6N M\u1ec0M QU\u1ea2N L\u00dd C\u0102NG-TIN', bold=True, size=16, space_after=4)
    add_centered(doc, 'T\u1ea0I C\u00c1C TR\u1ea0I T\u1ea0M GIAM, NH\u00c0 T\u1ea0M GI\u1eee', bold=True, size=14, space_after=4)
    add_centered(doc, 'C\u00d4NG AN C\u00c1C \u0110\u01a0N V\u1eca, \u0110\u1ecaA PH\u01af\u01a0NG', bold=True, size=14, space_after=24)

    for _ in range(4):
        doc.add_paragraph()

    add_centered(doc, 'H\u00c0 N\u1ed8I, N\u0102M 2026', bold=True, size=14)
    doc.add_page_break()

    # ===== TOC =====
    add_centered(doc, 'M\u1ee4C L\u1ee4C', bold=True, size=16, space_after=12)
    toc_items = [
        'I. Gi\u1edbi thi\u1ec7u v\u1ec1 ph\u1ea7n m\u1ec1m',
        'II. \u0110\u0103ng nh\u1eadp h\u1ec7 th\u1ed1ng',
        'III. B\u1ea3ng \u0111i\u1ec1u khi\u1ec3n t\u1ed5ng quan ',
        '1. Qu\u1ea3n l\u00fd \u0111\u01a1n h\u00e0ng',
        '   1.1. Danh s\u00e1ch \u0111\u01a1n h\u00e0ng',
        '   1.2. Phi\u1ebfu giao nh\u1eadn h\u00e0ng h\u00f3a',
        '   1.3. T\u1ed5ng h\u1ee3p b\u1ebfp',
        '2. Qu\u1ea7y thu ng\u00e2n',
        '   2.1. Tra c\u1ee9u th\u00f4ng tin can ph\u1ea1m nh\u00e2n',
        '   2.2. N\u1ea1p ti\u1ec1n t\u00e0i kho\u1ea3n l\u01b0u k\u00fd',
        '   2.3. Duy\u1ec7t \u0111\u01a1n h\u00e0ng ch\u1edd',
        '   2.4. \u0110\u1eb7t h\u00e0ng t\u1eeb th\u00e2n nh\u00e2n',
        '3. Qu\u1ea3n l\u00fd danh m\u1ee5c h\u00e0ng h\u00f3a',
        '4. Qu\u1ea3n l\u00fd t\u00e0i kho\u1ea3n l\u01b0u k\u00fd',
        '5. Qu\u00e9t phi\u1ebfu qua \u0111i\u1ec7n tho\u1ea1i',
        '6. Gi\u00e1m s\u00e1t qu\u00e9t',
        '7. Nh\u1eadt k\u00fd ho\u1ea1t \u0111\u1ed9ng',
        '8. Qu\u1ea3n l\u00fd bi\u1ec3u m\u1eabu OMR',
        '9. Nh\u1eadp \u0111\u01a1n h\u00e0ng th\u1ee7 c\u00f4ng',
        '10. Qu\u1ea3n l\u00fd c\u00e1n b\u1ed9',
        '11. C\u1ea5u h\u00ecnh thanh to\u00e1n',
        '12. C\u1ea5u h\u00ecnh h\u1ea1n m\u1ee9c mua h\u00e0ng',
        '13. \u0110\u1eb7t h\u00e0ng t\u1ef1 ph\u1ee5c v\u1ee5',
        '14. B\u00e1o c\u00e1o t\u00e0i ch\u00ednh',
        'IV. \u1ee8ng d\u1ee5ng m\u00e1y t\u00ednh (Electron)',
        'PH\u1ee4 L\u1ee4C',
    ]
    for item in toc_items:
        add_para(doc, item, size=12, space_after=2)
    doc.add_page_break()

    # ========================================================================
    # CONTENT
    # ========================================================================

    # I. Gioi thieu
    add_h1(doc, 'I. Gi\u1edbi thi\u1ec7u v\u1ec1 ph\u1ea7n m\u1ec1m')
    add_para(doc,
        'Ph\u1ea7n m\u1ec1m "Qu\u1ea3n l\u00fd C\u0103ng-tin Tr\u1ea1i t\u1ea1m giam" (Canteen Manager) '
        'l\u00e0 h\u1ec7 th\u1ed1ng qu\u1ea3n l\u00fd ho\u1ea1t \u0111\u1ed9ng mua b\u00e1n l\u01b0\u01a1ng th\u1ef1c, '
        'th\u1ef1c ph\u1ea9m, \u0111\u1ed3 d\u00f9ng sinh ho\u1ea1t t\u1ea1i c\u00e1c Tr\u1ea1i t\u1ea1m giam, '
        'Nh\u00e0 t\u1ea1m gi\u1eef C\u00f4ng an c\u00e1c \u0111\u01a1n v\u1ecb, \u0111\u1ecba ph\u01b0\u01a1ng. '
        'Ph\u1ea7n m\u1ec1m h\u1ed7 tr\u1ee3 to\u00e0n b\u1ed9 quy tr\u00ecnh t\u1eeb ti\u1ebfp nh\u1eadn \u0111\u01a1n h\u00e0ng '
        '(qua phi\u1ebfu qu\u00e9t OMR, kiosk, nh\u1eadp th\u1ee7 c\u00f4ng, \u0111\u1eb7t t\u1eeb th\u00e2n nh\u00e2n), '
        'qu\u1ea3n l\u00fd t\u00e0i kho\u1ea3n l\u01b0u k\u00fd, thanh to\u00e1n, in phi\u1ebfu giao nh\u1eadn, '
        '\u0111\u1ebfn t\u1ed5ng h\u1ee3p \u0111\u01a1n h\u00e0ng cho b\u1ebfp v\u00e0 b\u00e1o c\u00e1o th\u1ed1ng k\u00ea.',
        indent=0.5)

    add_para(doc, 'Ph\u1ea7n m\u1ec1m h\u1ed7 tr\u1ee3 02 h\u00ecnh th\u1ee9c tri\u1ec3n khai:', indent=0.5, bold=True)
    add_bullet(doc, 'H\u00ecnh th\u1ee9c 1 \u2013 M\u00e1y ch\u1ee7 (Docker): D\u00e0nh cho c\u01a1 s\u1edf c\u00f3 h\u1ea1 t\u1ea7ng m\u00e1y ch\u1ee7, s\u1eed d\u1ee5ng PostgreSQL, tri\u1ec3n khai qua Docker Compose.')
    add_bullet(doc, 'H\u00ecnh th\u1ee9c 2 \u2013 C\u00e0i \u0111\u1eb7t c\u00e1 nh\u00e2n (Electron): D\u00e0nh cho c\u01a1 s\u1edf nh\u1ecf, c\u00e0i \u0111\u1eb7t tr\u00ean m\u00e1y t\u00ednh Windows, s\u1eed d\u1ee5ng SQLite.')

    add_para(doc,
        'S\u1eed d\u1ee5ng c\u00e1c tr\u00ecnh duy\u1ec7t Chrome, Firefox, Edge (phi\u00ean b\u1ea3n m\u1edbi nh\u1ea5t) '
        '\u0111\u1ec3 truy c\u1eadp ph\u1ea7n m\u1ec1m theo \u0111\u1ecba ch\u1ec9 m\u1ea1ng n\u1ed9i b\u1ed9 do qu\u1ea3n tr\u1ecb vi\u00ean cung c\u1ea5p.',
        indent=0.5)

    add_para(doc, 'Y\u00eau c\u1ea7u h\u1ec7 th\u1ed1ng t\u1ed1i thi\u1ec3u:', indent=0.5, bold=True)
    add_bullet(doc, 'M\u00e1y ch\u1ee7: Ubuntu 20.04+, RAM 4GB, HDD 50GB, Docker Engine')
    add_bullet(doc, 'M\u00e1y c\u00e1 nh\u00e2n (Electron): Windows 10/11 64-bit, RAM 4GB, HDD 10GB')
    add_bullet(doc, 'M\u00e1y tr\u1ea1m: Tr\u00ecnh duy\u1ec7t Chrome/Firefox/Edge phi\u00ean b\u1ea3n m\u1edbi nh\u1ea5t')

    # ---- II. Dang nhap ----
    add_h1(doc, 'II. \u0110\u0103ng nh\u1eadp h\u1ec7 th\u1ed1ng')
    add_step(doc,
        'B\u01b0\u1edbc 1: M\u1edf tr\u00ecnh duy\u1ec7t web, nh\u1eadp \u0111\u1ecba ch\u1ec9 ph\u1ea7n m\u1ec1m do qu\u1ea3n tr\u1ecb vi\u00ean cung c\u1ea5p '
        '(v\u00ed d\u1ee5: http://192.168.1.100:3000).')
    add_step(doc,
        'B\u01b0\u1edbc 2: M\u00e0n h\u00ecnh hi\u1ec3n th\u1ecb giao di\u1ec7n \u0111\u0103ng nh\u1eadp. '
        'Anh/Ch\u1ecb nh\u1eadp t\u00ean \u0111\u0103ng nh\u1eadp v\u00e0 m\u1eadt kh\u1ea9u.')

    add_image(doc, 'login.png', 'H\u00ecnh 1: Giao di\u1ec7n \u0111\u0103ng nh\u1eadp h\u1ec7 th\u1ed1ng')

    add_step(doc, 'B\u01b0\u1edbc 3: Click n\u00fat "\u0110\u0103ng nh\u1eadp" \u0111\u1ec3 truy c\u1eadp h\u1ec7 th\u1ed1ng.')
    add_note(doc,
        'Phi\u00ean \u0111\u0103ng nh\u1eadp c\u00f3 hi\u1ec7u l\u1ef1c trong 12 gi\u1edd. '
        'T\u00f9y theo vai tr\u00f2 (Qu\u1ea3n tr\u1ecb vi\u00ean / C\u00e1n b\u1ed9 nghi\u1ec7p v\u1ee5 / Thu ng\u00e2n), giao di\u1ec7n s\u1ebd hi\u1ec3n th\u1ecb c\u00e1c ch\u1ee9c n\u0103ng t\u01b0\u01a1ng \u1ee9ng. '
        'C\u00e1n b\u1ed9 \u0111\u01b0\u1ee3c ph\u00e2n v\u00f9ng theo khu giam.')

    # ---- III. Dashboard ----
    add_h1(doc, 'III. B\u1ea3ng \u0111i\u1ec1u khi\u1ec3n t\u1ed5ng quan ')
    add_para(doc,
        'Sau khi \u0111\u0103ng nh\u1eadp th\u00e0nh c\u00f4ng, h\u1ec7 th\u1ed1ng hi\u1ec3n th\u1ecb m\u00e0n h\u00ecnh '
        'B\u1ea3ng \u0111i\u1ec1u khi\u1ec3n t\u1ed5ng quan v\u1edbi c\u00e1c th\u00f4ng tin ch\u00ednh:',
        indent=0.5)
    add_bullet(doc, 'T\u1ed5ng s\u1ed1 \u0111\u01a1n h\u00e0ng')
    add_bullet(doc, 'S\u1ed1 \u0111\u01a1n h\u00e0ng ch\u1edd duy\u1ec7t (Pending)')
    add_bullet(doc, 'S\u1ed1 \u0111\u01a1n h\u00e0ng \u0111\u00e3 thanh to\u00e1n (Paid)')
    add_bullet(doc, 'T\u1ed5ng doanh thu')

    add_image(doc, 'dashboard.png', 'H\u00ecnh 2: B\u1ea3ng \u0111i\u1ec1u khi\u1ec3n t\u1ed5ng quan ')

    add_para(doc,
        'Thanh \u0111i\u1ec1u h\u01b0\u1edbng b\u00ean tr\u00e1i hi\u1ec3n th\u1ecb c\u00e1c danh m\u1ee5c ch\u1ee9c n\u0103ng '
        't\u00f9y theo vai tr\u00f2:',
        indent=0.5)

    add_image(doc, 'app-shell.png', 'H\u00ecnh 3: Giao di\u1ec7n ch\u00ednh v\u1edbi thanh \u0111i\u1ec1u h\u01b0\u1edbng')

    # Role table
    role_table = doc.add_table(rows=1, cols=4)
    role_table.style = 'Table Grid'
    role_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = role_table.rows[0].cells
    for i, text in enumerate(['Ch\u1ee9c n\u0103ng', 'Qu\u1ea3n tr\u1ecb vi\u00ean', 'C\u00e1n b\u1ed9 NV', 'Thu ng\u00e2n']):
        hdr[i].paragraphs[0].text = ''
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_run_tnr(p, text, bold=True, size=10)

    role_data = [
        ('T\u1ed5ng quan', '\u2713', '\u2713', '\u2713'),
        ('\u0110\u01a1n h\u00e0ng', '\u2713', '\u2713', ''),
        ('Phi\u1ebfu giao nh\u1eadn', '\u2713', '', ''),
        ('T\u1ed5ng h\u1ee3p b\u1ebfp', '\u2713', '\u2713', ''),
        ('Qu\u1ea7y thu ng\u00e2n', '\u2713', '', '\u2713'),
        ('Qu\u1ea3n l\u00fd th\u1ef1c \u0111\u01a1n', '\u2713', '', ''),
        ('T\u00e0i kho\u1ea3n l\u01b0u k\u00fd', '\u2713', '', ''),
        ('Gi\u00e1m s\u00e1t qu\u00e9t', '\u2713', '\u2713', '\u2713'),
        ('Nh\u1eadt k\u00fd ho\u1ea1t \u0111\u1ed9ng', '\u2713', '', ''),
        ('B\u00e1o c\u00e1o t\u00e0i ch\u00ednh', '\u2713', '', ''),
        ('In bi\u1ec3u m\u1eabu OMR', '\u2713', '', ''),
        ('Nh\u1eadp \u0111\u01a1n th\u1ee7 c\u00f4ng', '\u2713', '\u2713', ''),
        ('Qu\u1ea3n l\u00fd c\u00e1n b\u1ed9', '\u2713', '', ''),
        ('C\u1ea5u h\u00ecnh thanh to\u00e1n', '\u2713', '', ''),
        ('C\u1ea5u h\u00ecnh h\u1ea1n m\u1ee9c', '\u2713', '', ''),
    ]
    for row_data in role_data:
        row = role_table.add_row().cells
        for j, val in enumerate(row_data):
            row[j].paragraphs[0].text = ''
            p = row[j].paragraphs[0]
            if j > 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_run_tnr(p, val, size=10)
    doc.add_paragraph()

    # ===== 1. QUAN LY DON HANG =====
    add_h2(doc, '1. Qu\u1ea3n l\u00fd \u0111\u01a1n h\u00e0ng')

    add_h3(doc, '1.1. Danh s\u00e1ch \u0111\u01a1n h\u00e0ng')
    add_step(doc,
        'B\u01b0\u1edbc 1: T\u1ea1i thanh menu, anh/ch\u1ecb ch\u1ecdn "\u0110\u01a1n h\u00e0ng". '
        'H\u1ec7 th\u1ed1ng hi\u1ec3n th\u1ecb danh s\u00e1ch \u0111\u01a1n h\u00e0ng.')
    add_step(doc,
        'B\u01b0\u1edbc 2: L\u1ecdc danh s\u00e1ch theo: kho\u1ea3ng th\u1eddi gian, tr\u1ea1ng th\u00e1i '
        '(\u0110ang ho\u1ea1t \u0111\u1ed9ng / \u0110\u00e3 thay th\u1ebf / \u0110\u00e3 t\u1eeb ch\u1ed1i), thanh to\u00e1n (\u0110\u00e3 thanh to\u00e1n / Ch\u01b0a thanh to\u00e1n), '
        'ngu\u1ed3n (qu\u00e9t phi\u1ebfu / m\u00e1y qu\u00e9t / th\u00e2n nh\u00e2n / th\u1ee7 c\u00f4ng).')

    add_image(doc, 'orders.png', 'H\u00ecnh 4: Danh s\u00e1ch \u0111\u01a1n h\u00e0ng')

    add_step(doc,
        'B\u01b0\u1edbc 3: Click v\u00e0o \u0111\u01a1n h\u00e0ng \u0111\u1ec3 xem chi ti\u1ebft: '
        'th\u00f4ng tin can ph\u1ea1m nh\u00e2n, danh s\u00e1ch m\u1eb7t h\u00e0ng, s\u1ed1 l\u01b0\u1ee3ng, '
        '\u0111\u01a1n gi\u00e1, t\u1ed5ng ti\u1ec1n, th\u00f4ng tin thanh to\u00e1n.')

    add_h3(doc, '1.2. Phi\u1ebfu giao nh\u1eadn h\u00e0ng h\u00f3a ')
    add_step(doc, 'B\u01b0\u1edbc 1: Ch\u1ecdn "Phi\u1ebfu giao nh\u1eadn". Ch\u1ecdn ng\u00e0y giao h\u00e0ng.')
    add_step(doc, 'B\u01b0\u1edbc 2: H\u1ec7 th\u1ed1ng t\u1ed5ng h\u1ee3p t\u1ea5t c\u1ea3 \u0111\u01a1n h\u00e0ng \u0111\u00e3 thanh to\u00e1n (PAID) c\u1ee7a ng\u00e0y \u0111\u00f3.')
    add_step(doc, 'B\u01b0\u1edbc 3: M\u1ed7i phi\u1ebfu g\u1ed3m: th\u00f4ng tin can ph\u1ea1m nh\u00e2n, danh s\u00e1ch m\u1eb7t h\u00e0ng, t\u1ed5ng ti\u1ec1n, s\u1ed1 d\u01b0 c\u00f2n l\u1ea1i.')

    add_image(doc, 'vouchers.png', 'H\u00ecnh 5: Phi\u1ebfu giao nh\u1eadn h\u00e0ng h\u00f3a (Vouchers)')

    add_step(doc, 'B\u01b0\u1edbc 4: Click "In phi\u1ebfu" \u0111\u1ec3 t\u1ea3i v\u1ec1 m\u00e1y t\u00ednh v\u00e0 in.')

    add_h3(doc, '1.3. T\u1ed5ng h\u1ee3p b\u1ebfp ')
    add_step(doc, 'B\u01b0\u1edbc 1: Ch\u1ecdn "T\u1ed5ng h\u1ee3p b\u1ebfp". Ch\u1ecdn ng\u00e0y giao h\u00e0ng.')
    add_step(doc, 'B\u01b0\u1edbc 2: H\u1ec7 th\u1ed1ng hi\u1ec3n th\u1ecb b\u1ea3ng t\u1ed5ng h\u1ee3p: t\u00ean m\u1eb7t h\u00e0ng, m\u00e3 h\u00e0ng, ph\u00e2n lo\u1ea1i, t\u1ed5ng s\u1ed1 l\u01b0\u1ee3ng, \u0111\u01a1n gi\u00e1, t\u1ed5ng th\u00e0nh ti\u1ec1n.')

    add_image(doc, 'kitchen-summary.png', 'H\u00ecnh 6: T\u1ed5ng h\u1ee3p b\u1ebfp ')

    add_step(doc, 'B\u01b0\u1edbc 3: In b\u1ea3ng t\u1ed5ng h\u1ee3p \u0111\u1ec3 chuy\u1ec3n cho b\u1ed9 ph\u1eadn b\u1ebfp.')

    # ===== 2. COUNTER =====
    add_h2(doc, '2. Qu\u1ea7y thu ng\u00e2n ')
    add_para(doc, 'D\u00e0nh cho c\u00e1n b\u1ed9 vai tr\u00f2 Thu ng\u00e2n ho\u1eb7c Qu\u1ea3n tr\u1ecb vi\u00ean.', indent=0.5)

    add_h3(doc, '2.1. Tra c\u1ee9u th\u00f4ng tin can ph\u1ea1m nh\u00e2n')
    add_step(doc, 'B\u01b0\u1edbc 1: Ch\u1ecdn "Qu\u1ea7y thu ng\u00e2n". Nh\u1eadp m\u00e3 l\u01b0u k\u00fd v\u00e0o \u00f4 t\u00ecm ki\u1ebfm.')

    add_image(doc, 'counter-empty.png', 'H\u00ecnh 7: Giao di\u1ec7n qu\u1ea7y thu ng\u00e2n - \u00f4 t\u00ecm ki\u1ebfm v\u00e0 h\u00e0ng \u0111\u1ee3i \u0111\u01a1n ch\u1edd duy\u1ec7t')

    add_step(doc,
        'B\u01b0\u1edbc 2: H\u1ec7 th\u1ed1ng hi\u1ec3n th\u1ecb: h\u1ecd t\u00ean, m\u00e3 l\u01b0u k\u00fd, '
        'bu\u1ed3ng giam, khu giam, s\u1ed1 d\u01b0 t\u00e0i kho\u1ea3n, h\u1ea1n m\u1ee9c chi ti\u00eau c\u00f2n l\u1ea1i, '
        'l\u1ecbch s\u1eed giao d\u1ecbch g\u1ea7n \u0111\u00e2y.')

    add_image(doc, 'counter-lookup.png', 'H\u00ecnh 8: Th\u00f4ng tin can ph\u1ea1m nh\u00e2n sau khi tra c\u1ee9u')

    add_h3(doc, '2.2. N\u1ea1p ti\u1ec1n t\u00e0i kho\u1ea3n l\u01b0u k\u00fd')
    add_step(doc, 'B\u01b0\u1edbc 1: Sau khi tra c\u1ee9u, click "N\u1ea1p ti\u1ec1n" .')
    add_step(doc, 'B\u01b0\u1edbc 2: Nh\u1eadp s\u1ed1 ti\u1ec1n n\u1ea1p (VND) v\u00e0 click "X\u00e1c nh\u1eadn".')

    add_image(doc, 'counter-prisoner-loaded.png',
              'H\u00ecnh 9: H\u1ed3 s\u01a1 ph\u1ea1m nh\u00e2n, form n\u1ea1p ti\u1ec1n v\u00e0 form \u0111\u1eb7t h\u00e0ng ng\u01b0\u1eddi th\u00e2n')

    add_step(doc,
        'B\u01b0\u1edbc 3: H\u1ec7 th\u1ed1ng ghi nh\u1eadn giao d\u1ecbch TOPUP v\u00e0o s\u1ed5 c\u00e1i, '
        'c\u1eadp nh\u1eadt s\u1ed1 d\u01b0 t\u1ef1 \u0111\u1ed9ng.')
    add_note(doc,
        'Giao d\u1ecbch n\u1ea1p ti\u1ec1n kh\u00f4ng th\u1ec3 s\u1eeda/x\u00f3a (m\u00f4 h\u00ecnh s\u1ed5 c\u00e1i b\u1ea5t bi\u1ebfn). '
        'N\u1ebfu n\u1ea1p sai, th\u1ef1c hi\u1ec7n giao d\u1ecbch ho\u00e0n ti\u1ec1n (REFUND) ri\u00eang.')

    add_h3(doc, '2.3. Duy\u1ec7t \u0111\u01a1n h\u00e0ng ch\u1edd')
    add_step(doc, 'B\u01b0\u1edbc 1: Ch\u1ecdn tab "\u0110\u01a1n h\u00e0ng ch\u1edd duy\u1ec7t" .')
    add_step(doc, 'B\u01b0\u1edbc 2: Ch\u1ecdn \u0111\u01a1n h\u00e0ng c\u1ea7n duy\u1ec7t, xem chi ti\u1ebft.')

    add_image(doc, 'counter-pending-queue.png',
              'H\u00ecnh 10: H\u00e0ng \u0111\u1ee3i \u0111\u01a1n h\u00e0ng ch\u1edd duy\u1ec7t t\u1ea1i qu\u1ea7y thu ng\u00e2n')

    add_step(doc, 'B\u01b0\u1edbc 3a \u2013 Duy\u1ec7t: Ch\u1ecdn ph\u01b0\u01a1ng th\u1ee9c thanh to\u00e1n:')
    add_bullet(doc, '"Tr\u1eeb s\u1ed1 d\u01b0" : T\u1ef1 \u0111\u1ed9ng tr\u1eeb t\u1eeb t\u00e0i kho\u1ea3n l\u01b0u k\u00fd.', indent_cm=1.5)
    add_bullet(doc, '"Ti\u1ec1n m\u1eb7t" : Ghi nh\u1eadn thanh to\u00e1n ti\u1ec1n m\u1eb7t.', indent_cm=1.5)
    add_bullet(doc, '"Chuy\u1ec3n kho\u1ea3n" : Nh\u1eadp m\u00e3 giao d\u1ecbch v\u00e0 s\u1ed1 ti\u1ec1n nh\u1eadn.', indent_cm=1.5)
    add_step(doc,
        'B\u01b0\u1edbc 3b \u2013 T\u1eeb ch\u1ed1i: Click "T\u1eeb ch\u1ed1i", nh\u1eadp l\u00fd do. '
        'H\u1ec7 th\u1ed1ng t\u1ef1 \u0111\u1ed9ng ho\u00e0n ti\u1ec1n n\u1ebfu \u0111\u00e3 tr\u1eeb tr\u01b0\u1edbc \u0111\u00f3.')

    add_h3(doc, '2.4. \u0110\u1eb7t h\u00e0ng t\u1eeb th\u00e2n nh\u00e2n')
    add_step(doc, 'B\u01b0\u1edbc 1: Click "\u0110\u1eb7t h\u00e0ng t\u1eeb th\u00e2n nh\u00e2n" .')
    add_step(doc, 'B\u01b0\u1edbc 2: Nh\u1eadp m\u00e3 l\u01b0u k\u00fd. Ch\u1ecdn m\u1eb7t h\u00e0ng v\u00e0 s\u1ed1 l\u01b0\u1ee3ng t\u1eeb danh m\u1ee5c.')
    add_step(doc, 'B\u01b0\u1edbc 3: Ch\u1ecdn ph\u01b0\u01a1ng th\u1ee9c thanh to\u00e1n (ti\u1ec1n m\u1eb7t/chuy\u1ec3n kho\u1ea3n).')
    add_step(doc, 'B\u01b0\u1edbc 4: Click "T\u1ea1o \u0111\u01a1n v\u00e0 thanh to\u00e1n". H\u1ec7 th\u1ed1ng t\u1ea1o \u0111\u01a1n h\u00e0ng ACTIVE + PAID.')
    add_note(doc,
        'Ti\u1ec1n kh\u00f4ng tr\u1eeb t\u1eeb t\u00e0i kho\u1ea3n l\u01b0u k\u00fd m\u00e0 do th\u00e2n nh\u00e2n thanh to\u00e1n tr\u1ef1c ti\u1ebfp.')

    # ===== 3. MENU =====
    add_h2(doc, '3. Qu\u1ea3n l\u00fd danh m\u1ee5c h\u00e0ng h\u00f3a ')
    add_para(doc, 'D\u00e0nh cho Qu\u1ea3n tr\u1ecb vi\u00ean (Qu\u1ea3n tr\u1ecb vi\u00ean).', indent=0.5)
    add_step(doc,
        'B\u01b0\u1edbc 1: Ch\u1ecdn "Qu\u1ea3n l\u00fd th\u1ef1c \u0111\u01a1n". H\u1ec7 th\u1ed1ng hi\u1ec3n th\u1ecb danh s\u00e1ch m\u1eb7t h\u00e0ng: '
        'm\u00e3 h\u00e0ng, t\u00ean, ph\u00e2n lo\u1ea1i (Th\u1ef1c ph\u1ea9m / \u0110\u1ed3 d\u00f9ng thi\u1ebft y\u1ebfu), gi\u00e1 b\u00e1n, tr\u1ea1ng th\u00e1i.')

    add_image(doc, 'menu-config.png', 'H\u00ecnh 11: Qu\u1ea3n l\u00fd danh m\u1ee5c h\u00e0ng h\u00f3a')

    add_step(doc, 'Th\u00eam m\u1edbi: Click "Th\u00eam m\u1edbi" \u2192 Nh\u1eadp t\u00ean, ph\u00e2n lo\u1ea1i, gi\u00e1 b\u00e1n \u2192 "L\u01b0u". M\u00e3 h\u00e0ng \u0111\u01b0\u1ee3c g\u00e1n t\u1ef1 \u0111\u1ed9ng.')
    add_step(doc, 'S\u1eeda: Click "S\u1eeda" tr\u00ean d\u00f2ng m\u1eb7t h\u00e0ng \u2192 Ch\u1ec9nh s\u1eeda \u2192 "C\u1eadp nh\u1eadt".')
    add_step(doc, 'V\u00f4 hi\u1ec7u h\u00f3a: Click "Ng\u1eebng b\u00e1n" \u2192 M\u1eb7t h\u00e0ng chuy\u1ec3n sang Ng\u1eebng ho\u1ea1t \u0111\u1ed9ng, kh\u00f4ng hi\u1ec3n th\u1ecb tr\u00ean kiosk.')
    add_note(doc, 'M\u00e3 h\u00e0ng kh\u00f4ng th\u1ec3 thay \u0111\u1ed5i v\u00e0 kh\u00f4ng bao gi\u1edd t\u00e1i s\u1eed d\u1ee5ng.')

    # ===== 4. ACCOUNTS =====
    add_h2(doc, '4. Qu\u1ea3n l\u00fd t\u00e0i kho\u1ea3n l\u01b0u k\u00fd ')
    add_step(doc, 'B\u01b0\u1edbc 1: Ch\u1ecdn "T\u00e0i kho\u1ea3n l\u01b0u k\u00fd". Nh\u1eadp m\u00e3 l\u01b0u k\u00fd ho\u1eb7c h\u1ecd t\u00ean \u0111\u1ec3 t\u00ecm ki\u1ebfm.')
    add_step(doc,
        'B\u01b0\u1edbc 2: H\u1ec7 th\u1ed1ng hi\u1ec3n th\u1ecb: th\u00f4ng tin can ph\u1ea1m nh\u00e2n, s\u1ed1 d\u01b0 hi\u1ec7n t\u1ea1i, '
        'l\u1ecbch s\u1eed giao d\u1ecbch (s\u1ed5 c\u00e1i).')

    add_image(doc, 'audit.png', 'H\u00ecnh 12: Ki\u1ec3m to\u00e1n s\u1ed1 d\u01b0 t\u00e0i kho\u1ea3n (Audit)')

    add_para(doc,
        'L\u1ecbch s\u1eed giao d\u1ecbch bao g\u1ed3m: ng\u00e0y gi\u1edd, lo\u1ea1i giao d\u1ecbch '
        '(N\u1ea1p ti\u1ec1n / Tr\u1eeb ti\u1ec1n \u0111\u01a1n h\u00e0ng / Ho\u00e0n ti\u1ec1n h\u1ee7y \u0111\u01a1n / Ho\u00e0n ti\u1ec1n th\u1ee7 c\u00f4ng), s\u1ed1 ti\u1ec1n, s\u1ed1 d\u01b0 sau giao d\u1ecbch, '
        'c\u00e1n b\u1ed9 th\u1ef1c hi\u1ec7n.',
        indent=0.5)
    add_note(doc,
        'S\u1ed5 c\u00e1i ho\u1ea1t \u0111\u1ed9ng theo m\u00f4 h\u00ecnh ch\u1ec9 th\u00eam (ch\u1ec9 ghi th\u00eam). '
        'Kh\u00f4ng giao d\u1ecbch n\u00e0o b\u1ecb s\u1eeda/x\u00f3a.')

    # ===== 5. PHONE SCAN =====
    add_h2(doc, '5. Qu\u00e9t phi\u1ebfu qua \u0111i\u1ec7n tho\u1ea1i (Phone Scan)')
    add_para(doc, 'Ch\u1ee9c n\u0103ng qu\u00e9t phi\u1ebfu OMR b\u1eb1ng camera \u0111i\u1ec7n tho\u1ea1i, kh\u00f4ng c\u1ea7n m\u00e1y qu\u00e9t chuy\u00ean d\u1ee5ng.', indent=0.5)
    add_step(doc, 'B\u01b0\u1edbc 1: Truy c\u1eadp \u0111\u1ecba ch\u1ec9 qu\u00e9t phi\u1ebfu (v\u00ed d\u1ee5: http://192.168.1.100:3000/scan) tr\u00ean \u0111i\u1ec7n tho\u1ea1i.')
    add_step(doc, 'B\u01b0\u1edbc 2: H\u1ec7 th\u1ed1ng m\u1edf camera thi\u1ebft b\u1ecb. \u0110\u1eb7t phi\u1ebfu OMR v\u00e0o khung h\u00ecnh.')
    add_step(doc, 'B\u01b0\u1edbc 3: Click "Ch\u1ee5p" \u0111\u1ec3 ch\u1ee5p \u1ea3nh phi\u1ebfu. Xem l\u1ea1i \u1ea3nh, "Ch\u1ee5p l\u1ea1i" n\u1ebfu c\u1ea7n ho\u1eb7c "G\u1eedi" \u0111\u1ec3 x\u1eed l\u00fd.')
    add_step(doc, 'B\u01b0\u1edbc 4: H\u1ec7 th\u1ed1ng x\u1eed l\u00fd: \u0111\u1ecdc m\u00e3 QR, nh\u1eadn d\u1ea1ng v\u00f9ng \u0111\u00e1nh d\u1ea5u, \u0111\u1ecdc \u00f4 s\u1ed1 l\u01b0\u1ee3ng, \u0111\u1ed1i chi\u1ebfu danh m\u1ee5c.')
    add_step(doc, 'B\u01b0\u1edbc 5: K\u1ebft qu\u1ea3: "\u0110\u00e3 t\u1ea1o \u0111\u01a1n" (t\u1ef1 \u0111\u1ed9ng) / "C\u1ea7n xem x\u00e9t" (c\u00f3 c\u1ea3nh b\u00e1o) / "Kh\u00f4ng t\u00ecm th\u1ea5y".')
    add_step(doc, 'B\u01b0\u1edbc 6: Click "Qu\u00e9t ti\u1ebfp" \u0111\u1ec3 qu\u00e9t phi\u1ebfu ti\u1ebfp theo.')
    add_note(doc, 'Kh\u00f4ng c\u1ea7n \u0111\u0103ng nh\u1eadp. \u0110\u1ea3m b\u1ea3o \u00e1nh s\u00e1ng \u0111\u1ee7 v\u00e0 phi\u1ebfu kh\u00f4ng nh\u0103n.')

    # ===== 6. SCAN MONITOR =====
    add_h2(doc, '6. Gi\u00e1m s\u00e1t qu\u00e9t (Scan Monitor)')
    add_step(doc, 'B\u01b0\u1edbc 1: Ch\u1ecdn "Gi\u00e1m s\u00e1t qu\u00e9t". H\u1ec7 th\u1ed1ng hi\u1ec3n th\u1ecb b\u1ea3ng \u0111i\u1ec1u khi\u1ec3n:')
    add_bullet(doc, 'T\u1ed5ng l\u01b0\u1ee3t qu\u00e9t: S\u1ed1 l\u01b0\u1ee3t qu\u00e9t phi\u1ebfu qua \u0111i\u1ec7n tho\u1ea1i')
    add_bullet(doc, '\u0110\u01a1n h\u00e0ng t\u1ea1o: S\u1ed1 \u0111\u01a1n h\u00e0ng \u0111\u01b0\u1ee3c t\u1ea1o t\u1eeb qu\u00e9t phi\u1ebfu')
    add_bullet(doc, '\u0110\u01a1n \u0111\u00e3 thanh to\u00e1n: S\u1ed1 \u0111\u01a1n \u0111\u00e3 \u0111\u01b0\u1ee3c thanh to\u00e1n')
    add_bullet(doc, 'T\u1ed5ng doanh thu: T\u1ed5ng gi\u00e1 tr\u1ecb \u0111\u01a1n h\u00e0ng t\u1eeb qu\u00e9t phi\u1ebfu')
    add_step(doc, 'B\u01b0\u1edbc 2: B\u1ea3ng l\u1ecbch s\u1eed qu\u00e9t: m\u00e3 \u0111\u01a1n, ng\u00e0y ph\u1ee5c v\u1ee5, s\u1ed1 ti\u1ec1n, tr\u1ea1ng th\u00e1i, th\u1eddi gian.')
    add_step(doc, 'B\u01b0\u1edbc 3: Click "M\u1edf qu\u00e9t \u0111i\u1ec7n tho\u1ea1i" \u0111\u1ec3 m\u1edf trang qu\u00e9t trong tab m\u1edbi.')
    add_note(doc, 'D\u1eef li\u1ec7u t\u1ef1 \u0111\u1ed9ng l\u00e0m m\u1edbi m\u1ed7i 5 gi\u00e2y. C\u00f3 th\u1ec3 t\u1ea1m d\u1eebng/ti\u1ebfp t\u1ee5c.')

    # ===== 7. AUDIT LOG =====
    add_h2(doc, '7. Nh\u1eadt k\u00fd ho\u1ea1t \u0111\u1ed9ng (Audit Log)')
    add_step(doc, 'B\u01b0\u1edbc 1: Ch\u1ecdn "Nh\u1eadt k\u00fd ho\u1ea1t \u0111\u1ed9ng". H\u1ec7 th\u1ed1ng hi\u1ec3n th\u1ecb danh s\u00e1ch ho\u1ea1t \u0111\u1ed9ng g\u1ea7n \u0111\u00e2y.')
    add_step(doc, 'B\u01b0\u1edbc 2: M\u1ed7i b\u1ea3n ghi g\u1ed3m: th\u1eddi gian, c\u00e1n b\u1ed9, lo\u1ea1i thao t\u00e1c, \u0111\u1ed1i t\u01b0\u1ee3ng, chi ti\u1ebft thay \u0111\u1ed5i.')
    add_step(doc, 'B\u01b0\u1edbc 3: S\u1eed d\u1ee5ng b\u1ed9 l\u1ecdc \u0111\u1ec3 t\u00ecm ki\u1ebfm theo th\u1eddi gian, c\u00e1n b\u1ed9, lo\u1ea1i thao t\u00e1c.')
    add_note(doc, 'Nh\u1eadt k\u00fd kh\u00f4ng th\u1ec3 s\u1eeda/x\u00f3a (b\u1ea5t bi\u1ebfn). Ch\u1ec9 Qu\u1ea3n tr\u1ecb vi\u00ean c\u00f3 quy\u1ec1n xem.')

    # ===== 8. FORM PRINT =====
    add_h2(doc, '8. Qu\u1ea3n l\u00fd bi\u1ec3u m\u1eabu OMR ')
    add_step(doc, 'B\u01b0\u1edbc 1: Ch\u1ecdn "In bi\u1ec3u m\u1eabu". Ch\u1ecdn m\u1eabu phi\u1ebfu:')
    add_bullet(doc, 'Ch\u1ebf \u0111\u1ed9: CODE (ch\u1ec9 m\u00e3 h\u00e0ng) ho\u1eb7c FULL_LIST (m\u00e3 + t\u00ean h\u00e0ng)')
    add_bullet(doc, 'Kh\u1ed5 gi\u1ea5y: A4 / A5')
    add_bullet(doc, 'H\u01b0\u1edbng: D\u1ecdc (Portrait) / Ngang (Landscape)')
    add_step(doc, 'B\u01b0\u1edbc 2: Ch\u1ecdn can ph\u1ea1m nh\u00e2n c\u1ea7n in phi\u1ebfu (c\u00f3 th\u1ec3 ch\u1ecdn nhi\u1ec1u).')
    add_step(doc, 'B\u01b0\u1edbc 3: Click "In phi\u1ebfu". H\u1ec7 th\u1ed1ng t\u1ea1o phi\u1ebfu OMR v\u00e0 ghi nh\u1eadn phi\u1ebfu \u0111\u00e3 ph\u00e1t h\u00e0nh.')
    add_note(doc, 'Khi danh m\u1ee5c h\u00e0ng h\u00f3a thay \u0111\u1ed5i, h\u1ec7 th\u1ed1ng c\u1ea3nh b\u00e1o c\u1ea7n t\u1ea1o phi\u00ean b\u1ea3n bi\u1ec3u m\u1eabu m\u1edbi.')

    # ===== 9. ORDER FORM =====
    add_h2(doc, '9. Nh\u1eadp \u0111\u01a1n h\u00e0ng th\u1ee7 c\u00f4ng ')
    add_step(doc, 'B\u01b0\u1edbc 1: Ch\u1ecdn "Nh\u1eadp \u0111\u01a1n th\u1ee7 c\u00f4ng". Nh\u1eadp m\u00e3 l\u01b0u k\u00fd.')
    add_step(doc, 'B\u01b0\u1edbc 2: Ch\u1ecdn ng\u00e0y giao h\u00e0ng .')
    add_step(doc, 'B\u01b0\u1edbc 3: Ch\u1ecdn m\u1eb7t h\u00e0ng, nh\u1eadp s\u1ed1 l\u01b0\u1ee3ng. T\u1ed5ng ti\u1ec1n t\u1ef1 \u0111\u1ed9ng t\u00ednh.')
    add_step(doc, 'B\u01b0\u1edbc 4: Click "T\u1ea1o \u0111\u01a1n h\u00e0ng". \u0110\u01a1n h\u00e0ng ACTIVE + UNPAID, ch\u1edd duy\u1ec7t t\u1ea1i Counter.')
    add_note(doc,
        'N\u1ebfu \u0111\u00e3 c\u00f3 \u0111\u01a1n c\u00f9ng ng\u00e0y/ngu\u1ed3n, h\u1ec7 th\u1ed1ng cho ph\u00e9p thay th\u1ebf '
        '(\u0111\u01a1n c\u0169 \u2192 SUPERSEDED).')

    # ===== 10. OPERATORS =====
    add_h2(doc, '10. Qu\u1ea3n l\u00fd c\u00e1n b\u1ed9 ')
    add_step(doc,
        'B\u01b0\u1edbc 1: Ch\u1ecdn "Qu\u1ea3n l\u00fd c\u00e1n b\u1ed9". Hi\u1ec3n th\u1ecb danh s\u00e1ch c\u00e1n b\u1ed9: '
        't\u00ean \u0111\u0103ng nh\u1eadp, h\u1ecd t\u00ean, vai tr\u00f2, khu giam, tr\u1ea1ng th\u00e1i.')

    add_image(doc, 'operators.png', 'H\u00ecnh 13: Qu\u1ea3n l\u00fd c\u00e1n b\u1ed9')

    add_step(doc,
        'Th\u00eam m\u1edbi: Click "Th\u00eam m\u1edbi" \u2192 Nh\u1eadp t\u00ean \u0111\u0103ng nh\u1eadp, h\u1ecd t\u00ean, '
        'm\u1eadt kh\u1ea9u (\u22658 k\u00fd t\u1ef1), vai tr\u00f2, khu giam \u2192 "L\u01b0u".')
    add_step(doc, 'S\u1eeda: Click "S\u1eeda" \u2192 Ch\u1ec9nh s\u1eeda \u2192 "C\u1eadp nh\u1eadt".')
    add_step(doc, 'Kh\u00f3a: Click "Kh\u00f3a" \u2192 T\u00e0i kho\u1ea3n chuy\u1ec3n Ng\u1eebng ho\u1ea1t \u0111\u1ed9ng, kh\u00f4ng th\u1ec3 \u0111\u0103ng nh\u1eadp.')

    # ===== 11. PAYMENT CONFIG =====
    add_h2(doc, '11. C\u1ea5u h\u00ecnh thanh to\u00e1n ')
    add_step(doc, 'B\u01b0\u1edbc 1: Ch\u1ecdn "C\u1ea5u h\u00ecnh thanh to\u00e1n".')
    add_step(doc, 'B\u01b0\u1edbc 2: C\u1ea5u h\u00ecnh:')
    add_bullet(doc, 'Cho ph\u00e9p chuy\u1ec3n kho\u1ea3n ng\u00e2n h\u00e0ng: B\u1eadt/T\u1eaft')
    add_bullet(doc, 'Cho ph\u00e9p thanh to\u00e1n ti\u1ec1n m\u1eb7t: B\u1eadt/T\u1eaft')

    add_image(doc, 'payment-config.png', 'H\u00ecnh 14: C\u1ea5u h\u00ecnh thanh to\u00e1n')

    add_step(doc, 'B\u01b0\u1edbc 3: Click "L\u01b0u" \u0111\u1ec3 \u00e1p d\u1ee5ng.')

    # ===== 12. PURCHASE LIMIT =====
    add_h2(doc, '12. C\u1ea5u h\u00ecnh h\u1ea1n m\u1ee9c mua h\u00e0ng ')
    add_step(doc, 'B\u01b0\u1edbc 1: Ch\u1ecdn "C\u1ea5u h\u00ecnh h\u1ea1n m\u1ee9c".')
    add_step(doc, 'B\u01b0\u1edbc 2: Nh\u1eadp h\u1ea1n m\u1ee9c chi ti\u00eau h\u00e0ng th\u00e1ng (VND).')
    add_step(doc, 'B\u01b0\u1edbc 3: "L\u01b0u". H\u1ec7 th\u1ed1ng c\u1ea3nh b\u00e1o khi v\u01b0\u1ee3t h\u1ea1n m\u1ee9c t\u1ea1i Counter.')
    add_note(doc, 'H\u1ea1n m\u1ee9c t\u00ednh theo th\u00e1ng d\u01b0\u01a1ng l\u1ecbch, t\u1ef1 \u0111\u1ed9ng reset \u0111\u1ea7u th\u00e1ng.')

    # ===== 13. KIOSK =====
    add_h2(doc, '13. \u0110\u1eb7t h\u00e0ng t\u1ef1 ph\u1ee5c v\u1ee5 \u2013 \u0110\u1eb7t h\u00e0ng t\u1ef1 ph\u1ee5c v\u1ee5')
    add_para(doc,
        'Giao di\u1ec7n kh\u00f4ng c\u1ea7n \u0111\u0103ng nh\u1eadp, ph\u00f9 h\u1ee3p m\u00e1y t\u00ednh b\u1ea3ng '
        'ho\u1eb7c m\u00e0n h\u00ecnh c\u1ea3m \u1ee9ng.',
        indent=0.5)
    add_step(doc,
        'B\u01b0\u1edbc 1: Truy c\u1eadp \u0111\u1ecba ch\u1ec9 kiosk (v\u00ed d\u1ee5: http://192.168.1.100:3000/canteen).')

    add_image(doc, 'kiosk-entry.png', 'H\u00ecnh 15: \u0110\u1eb7t h\u00e0ng t\u1ef1 ph\u1ee5c v\u1ee5 \u2013 M\u00e0n h\u00ecnh nh\u1eadp m\u00e3 l\u01b0u k\u00fd')

    add_step(doc, 'B\u01b0\u1edbc 2: Nh\u1eadp m\u00e3 l\u01b0u k\u00fd qua b\u00e0n ph\u00edm s\u1ed1.')

    add_image(doc, 'kiosk-keypad-filled.png', 'H\u00ecnh 16: \u0110\u1eb7t h\u00e0ng t\u1ef1 ph\u1ee5c v\u1ee5 \u2013 \u0110\u00e3 nh\u1eadp m\u00e3 l\u01b0u k\u00fd')

    add_step(doc, 'B\u01b0\u1edbc 3: Ch\u1ecdn m\u1eb7t h\u00e0ng v\u00e0 s\u1ed1 l\u01b0\u1ee3ng.')

    add_image(doc, 'kiosk-menu.png', 'H\u00ecnh 17: \u0110\u1eb7t h\u00e0ng t\u1ef1 ph\u1ee5c v\u1ee5 \u2013 Ch\u1ecdn m\u00f3n \u0111\u1eb7t h\u00e0ng')

    add_step(doc,
        'B\u01b0\u1edbc 4: X\u00e1c nh\u1eadn \u0111\u01a1n h\u00e0ng. H\u1ec7 th\u1ed1ng t\u1ea1o \u0111\u01a1n PENDING, '
        'ch\u1edd duy\u1ec7t t\u1ea1i Counter.')

    # ===== 14. FINANCIAL REPORT =====
    add_h2(doc, '14. B\u00e1o c\u00e1o t\u00e0i ch\u00ednh (Financial Report)')
    add_step(doc, 'B\u01b0\u1edbc 1: Ch\u1ecdn "B\u00e1o c\u00e1o t\u00e0i ch\u00ednh". Ch\u1ecdn kho\u1ea3ng th\u1eddi gian c\u1ea7n b\u00e1o c\u00e1o.')
    add_step(doc, 'B\u01b0\u1edbc 2: H\u1ec7 th\u1ed1ng hi\u1ec3n th\u1ecb: t\u1ed5ng doanh thu, ph\u00e2n t\u00edch theo ngu\u1ed3n \u0111\u1eb7t h\u00e0ng, ph\u01b0\u01a1ng th\u1ee9c thanh to\u00e1n.')
    add_step(doc, 'B\u01b0\u1edbc 3: S\u1ed1 l\u01b0\u1ee3ng \u0111\u01a1n h\u00e0ng theo tr\u1ea1ng th\u00e1i.')
    add_note(doc, 'Ch\u1ec9 Qu\u1ea3n tr\u1ecb vi\u00ean c\u00f3 quy\u1ec1n xem b\u00e1o c\u00e1o t\u00e0i ch\u00ednh.')

    # ===== IV. ELECTRON =====
    add_h1(doc, 'IV. \u1ee8ng d\u1ee5ng m\u00e1y t\u00ednh (Electron)')
    add_para(doc,
        'D\u00e0nh cho c\u01a1 s\u1edf nh\u1ecf, c\u00e0i \u0111\u1eb7t tr\u00ean m\u00e1y t\u00ednh Windows '
        'kh\u00f4ng c\u1ea7n h\u1ea1 t\u1ea7ng m\u00e1y ch\u1ee7.',
        indent=0.5)

    add_h3(doc, '1. C\u00e0i \u0111\u1eb7t')
    add_step(doc, 'Ch\u1ea1y file "Canteen Manager Setup.exe" tr\u00ean Windows 10/11 64-bit. L\u00e0m theo h\u01b0\u1edbng d\u1eabn.')

    add_h3(doc, '2. Thi\u1ebft l\u1eadp l\u1ea7n \u0111\u1ea7u ')
    add_step(doc,
        'B\u01b0\u1edbc 1: Nh\u1eadp t\u00e0i kho\u1ea3n qu\u1ea3n tr\u1ecb vi\u00ean '
        '(m\u1eb7c \u0111\u1ecbnh: admin/admin123 \u2013 khuy\u1ebfn ngh\u1ecb \u0111\u1ed5i ngay).')
    add_step(doc, 'B\u01b0\u1edbc 2: C\u1ea5u h\u00ecnh m\u00fai gi\u1edd (m\u1eb7c \u0111\u1ecbnh: Asia/Saigon).')
    add_step(doc,
        'B\u01b0\u1edbc 3: H\u1ec7 th\u1ed1ng t\u1ef1 \u0111\u1ed9ng t\u1ea1o kh\u00f3a JWT, '
        'CSDL SQLite, ch\u1ea1y migration, kh\u1edfi \u0111\u1ed9ng backend.')
    add_step(doc, 'B\u01b0\u1edbc 4: Tr\u00ecnh duy\u1ec7t t\u1ef1 \u0111\u1ed9ng m\u1edf giao di\u1ec7n ph\u1ea7n m\u1ec1m.')

    add_h3(doc, '3. S\u1eed d\u1ee5ng h\u00e0ng ng\u00e0y')
    add_bullet(doc, '\u1ee8ng d\u1ee5ng ch\u1ea1y \u1edf khay h\u1ec7 th\u1ed1ng . Click \u0111\u00fap \u0111\u1ec3 m\u1edf.')
    add_bullet(doc, 'Backend t\u1ef1 \u0111\u1ed9ng kh\u1edfi \u0111\u1ed9ng/d\u1eebng theo \u1ee9ng d\u1ee5ng.')
    add_bullet(doc, 'C\u00e1c m\u00e1y trong LAN truy c\u1eadp qua \u0111\u1ecba ch\u1ec9 IP hi\u1ec3n th\u1ecb tr\u00ean giao di\u1ec7n.')
    add_bullet(doc, 'D\u1eef li\u1ec7u l\u01b0u t\u1ea1i: %APPDATA%/Canteen Manager/')
    add_note(doc, 'Khuy\u1ebfn ngh\u1ecb sao l\u01b0u th\u01b0 m\u1ee5c d\u1eef li\u1ec7u \u0111\u1ecbnh k\u1ef3.')

    # ===== PHU LUC =====
    doc.add_page_break()
    add_centered(doc, 'PH\u1ee4 L\u1ee4C', bold=True, size=16, space_after=12)

    # A. Trang thai don hang
    add_h3(doc, 'A. B\u1ea3ng tr\u1ea1ng th\u00e1i \u0111\u01a1n h\u00e0ng')
    t1 = doc.add_table(rows=1, cols=2)
    t1.style = 'Table Grid'
    t1.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t1.rows[0].cells
    hdr[0].paragraphs[0].text = ''
    add_run_tnr(hdr[0].paragraphs[0], 'Tr\u1ea1ng th\u00e1i', bold=True, size=11)
    hdr[1].paragraphs[0].text = ''
    add_run_tnr(hdr[1].paragraphs[0], 'M\u00f4 t\u1ea3', bold=True, size=11)

    for status, desc in [
        ('\u0110ang ho\u1ea1t \u0111\u1ed9ng', '\u0110\u01a1n h\u00e0ng \u0111ang ho\u1ea1t \u0111\u1ed9ng'),
        ('\u0110\u00e3 thay th\u1ebf', '\u0110\u00e3 b\u1ecb thay th\u1ebf b\u1edfi \u0111\u01a1n m\u1edbi (qu\u00e9t l\u1ea1i)'),
        ('\u0110\u00e3 t\u1eeb ch\u1ed1i', '\u0110\u00e3 b\u1ecb t\u1eeb ch\u1ed1i b\u1edfi c\u00e1n b\u1ed9 thu ng\u00e2n'),
        ('\u0110\u00e3 thanh to\u00e1n', '\u0110\u01a1n h\u00e0ng \u0111\u00e3 \u0111\u01b0\u1ee3c thanh to\u00e1n'),
        ('Ch\u01b0a thanh to\u00e1n', '\u0110ang ch\u1edd duy\u1ec7t t\u1ea1i qu\u1ea7y thu ng\u00e2n'),
    ]:
        add_table_row(t1, [status, desc])
    doc.add_paragraph()

    # B. Loai giao dich
    add_h3(doc, 'B. B\u1ea3ng lo\u1ea1i giao d\u1ecbch t\u00e0i kho\u1ea3n l\u01b0u k\u00fd')
    t2 = doc.add_table(rows=1, cols=2)
    t2.style = 'Table Grid'
    t2.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t2.rows[0].cells
    hdr[0].paragraphs[0].text = ''
    add_run_tnr(hdr[0].paragraphs[0], 'Lo\u1ea1i giao d\u1ecbch', bold=True, size=11)
    hdr[1].paragraphs[0].text = ''
    add_run_tnr(hdr[1].paragraphs[0], 'M\u00f4 t\u1ea3', bold=True, size=11)

    for txn, desc in [
        ('N\u1ea1p ti\u1ec1n', 'Ti\u1ec1n th\u00e2n nh\u00e2n g\u1eedi \u0111\u01b0\u1ee3c n\u1ea1p v\u00e0o t\u00e0i kho\u1ea3n l\u01b0u k\u00fd'),
        ('Tr\u1eeb ti\u1ec1n \u0111\u01a1n h\u00e0ng', 'Tr\u1eeb ti\u1ec1n khi thanh to\u00e1n \u0111\u01a1n h\u00e0ng'),
        ('Ho\u00e0n ti\u1ec1n h\u1ee7y \u0111\u01a1n', 'Ho\u00e0n ti\u1ec1n khi \u0111\u01a1n h\u00e0ng b\u1ecb h\u1ee7y ho\u1eb7c thay th\u1ebf'),
        ('Ho\u00e0n ti\u1ec1n th\u1ee7 c\u00f4ng', 'C\u00e1n b\u1ed9 th\u1ef1c hi\u1ec7n ho\u00e0n ti\u1ec1n th\u1ee7 c\u00f4ng'),
    ]:
        add_table_row(t2, [txn, desc])
    doc.add_paragraph()

    # C. Trang thai ket qua quet phieu
    add_h3(doc, 'C. B\u1ea3ng tr\u1ea1ng th\u00e1i k\u1ebft qu\u1ea3 qu\u00e9t phi\u1ebfu')
    t3 = doc.add_table(rows=1, cols=2)
    t3.style = 'Table Grid'
    t3.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t3.rows[0].cells
    hdr[0].paragraphs[0].text = ''
    add_run_tnr(hdr[0].paragraphs[0], 'Tr\u1ea1ng th\u00e1i', bold=True, size=11)
    hdr[1].paragraphs[0].text = ''
    add_run_tnr(hdr[1].paragraphs[0], 'M\u00f4 t\u1ea3', bold=True, size=11)

    for scan, desc in [
        ('\u0110\u00e3 t\u1ea1o \u0111\u01a1n', '\u0110\u01a1n h\u00e0ng \u0111\u01b0\u1ee3c t\u1ea1o t\u1ef1 \u0111\u1ed9ng (\u0111\u1ed9 tin c\u1eady cao)'),
        ('C\u1ea7n xem x\u00e9t', 'C\u1ea7n ki\u1ec3m tra l\u1ea1i (c\u00f3 c\u1ea3nh b\u00e1o ho\u1eb7c \u0111\u1ed9 tin c\u1eady th\u1ea5p)'),
        ('Kh\u00f4ng t\u00ecm th\u1ea5y', 'Kh\u00f4ng nh\u1eadn d\u1ea1ng \u0111\u01b0\u1ee3c m\u1eb7t h\u00e0ng n\u00e0o'),
    ]:
        add_table_row(t3, [scan, desc])

    # Save
    output = 'pitch-docs/02-HDSD-Phan-mem-Canteen.docx'
    doc.save(output)
    print(f'Document 2 saved: {output}')
    print(f'Embedded 17 screenshots from tools/screenshots/')


if __name__ == '__main__':
    main()
